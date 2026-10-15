"""
Bulk product import: reading Excel / CSV files, mapping rows to products, and the template
"""
import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Dict, List, Optional, Tuple

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError

from app.core.batches import capitalize_words, display_product_name, STORE_BRAND
from app.models.product import Product

MAX_IMPORT_ROWS = 5000
MAX_IMPORT_BYTES = 5 * 1024 * 1024
LOOSE_UNIT = "kg"
IMPORT_REORDER_LEVEL = Decimal("10")

# field, template header, required, description, extra header names accepted
COLUMNS = [
    ("item_name", "Item Name", True, "Name without size or brand, e.g. Salt", ("item", "name")),
    ("product_type", "Product Type", True, "Manufacturer, Own packed or Loose (sold per kg)", ("type",)),
    ("brand", "Brand", False, "Manufacturer items only (Own packed get MF, Loose none)", ()),
    ("net_quantity", "Pack Size", False, "Quantity in the pack, e.g. 1 or 500", ("net quantity",)),
    ("net_unit", "Pack Unit", False, "g, kg, ml, L, pcs or pages", ("net unit",)),
    ("hsn_code", "HSN Code", False, "HSN code", ("hsn",)),
    ("tax_percent", "GST %", False, "GST rate, e.g. 5", ("gst percent", "tax percent", "tax")),
    ("mrp", "MRP", False, "Maximum retail price", ()),
    ("selling_price", "Selling Price", False, "Cannot be more than MRP", ("price", "sale price")),
    ("expiry_date", "Expiry Date", False, "Date cell, or text DD/MM/YYYY (e.g. 30/03/2027)", ("expiry",)),
]
LABELS = {f: header for f, header, *_ in COLUMNS}
EXAMPLES = [
    ["Salt", "Manufacturer", "Tata", 1, "kg", None, None, 28, 27, "30/03/2027"],
    ["Toor Dal", "Own packed", None, 500, "g", None, None, 90, 85, "15/01/2027"],
    ["Sugar", "Loose", None, None, None, None, None, None, 45, None],
]
PRODUCT_TYPES = {"manufacturer": Product.TYPE_MANUFACTURER, "ownpacked": Product.TYPE_OWN_PACKED,
                 "loose": Product.TYPE_LOOSE}
UNITS = {"g": "g", "gm": "g", "gms": "g", "gram": "g", "grams": "g", "kg": "kg", "kgs": "kg",
         "ml": "ml", "l": "L", "ltr": "L", "litre": "L", "liter": "L",
         "pc": "pcs", "pcs": "pcs", "page": "pages", "pages": "pages"}
DECIMAL_FIELDS = ("net_quantity", "tax_percent", "mrp", "selling_price")
TEXT_FIELDS = ("brand", "hsn_code")


class ImportRowError(Exception):
    """A row that cannot be imported; the message lists what is wrong"""


def normalize_header(value: Any) -> str:
    text = re.sub(r"\(.*?\)", " ", str(value or "").lower())
    return " ".join(re.sub(r"[^a-z0-9]+", " ", text).split())


HEADER_FIELDS = {}
for _field, _header, _req, _desc, _aliases in COLUMNS:
    for _name in (_field, _header, *_aliases):
        HEADER_FIELDS[normalize_header(_name)] = _field


def cell_text(value: Any) -> Optional[str]:
    """Cell value as text; whole numbers from Excel lose their trailing .0"""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    if re.fullmatch(r"[0-9]+\.0+", text):
        text = text.split(".")[0]
    return text or None


def read_import_file(filename: str, content: bytes) -> Tuple[List[Any], List[Tuple[int, List[Any]]]]:
    """Header row and data rows (with their sheet row numbers) of an .xlsx or .csv file"""
    name = filename.lower()
    if name.endswith(".xls"):
        raise HTTPException(status_code=400, detail="Old .xls files are not supported; save the file as .xlsx")
    if name.endswith((".xlsx", ".xlsm")):
        if not content.startswith(b"PK"):
            raise HTTPException(status_code=400, detail="The file is not a valid Excel (.xlsx) file")
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = workbook["Products"] if "Products" in workbook.sheetnames else workbook.worksheets[0]
            raw = [list(r) for r in sheet.iter_rows(values_only=True)]
            workbook.close()
        except Exception:
            raise HTTPException(status_code=400, detail="Could not read the Excel file")
    elif name.endswith(".csv"):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("cp1252", errors="replace")
        raw = [[v if v.strip() else None for v in r] for r in csv.reader(io.StringIO(text))]
    else:
        raise HTTPException(status_code=400, detail="Upload an Excel (.xlsx) or CSV file")

    numbered = [(i, r) for i, r in enumerate(raw, start=1)
                if any(v is not None and str(v).strip() for v in r)]
    if not numbered:
        raise HTTPException(status_code=400, detail="The file is empty")
    _, headers = numbered[0]
    while headers and (headers[-1] is None or not str(headers[-1]).strip()):
        headers = headers[:-1]
    rows = [(i, (list(r) + [None] * len(headers))[:len(headers)]) for i, r in numbered[1:]]
    if not rows:
        raise HTTPException(status_code=400, detail="The file has no product rows")
    return headers, rows


def map_headers(headers: List[Any]) -> Dict[str, int]:
    """Column index of each known field; Item Name and Product Type columns are required"""
    columns: Dict[str, int] = {}
    for index, header in enumerate(headers):
        name = HEADER_FIELDS.get(normalize_header(header))
        if name and name not in columns:
            columns[name] = index
    missing = [LABELS[f] for f, _h, required, *_ in COLUMNS if required and f not in columns]
    if missing:
        raise HTTPException(status_code=400,
                            detail=f"Missing column(s): {', '.join(missing)}. Use the import template.")
    return columns


def import_display_name(product_data) -> str:
    """Display name the product will get; mirrors apply_naming for a new product"""
    product_type = product_data.product_type
    item_name = capitalize_words((product_data.item_name or "").strip())
    if product_type == Product.TYPE_OWN_PACKED:
        brand = STORE_BRAND
    elif product_type == Product.TYPE_LOOSE:
        brand = None
    else:
        brand = capitalize_words((product_data.brand or "").strip()) or None
    return display_product_name(product_type, item_name, product_data.net_quantity,
                                product_data.net_unit, brand) or item_name


def parse_decimal(value: Any, label: str) -> Decimal:
    text = (cell_text(value) or "").replace(",", "").replace("₹", "").strip()
    try:
        number = Decimal(text)
    except InvalidOperation:
        raise ImportRowError(f"{label} must be a number")
    if not number.is_finite():
        raise ImportRowError(f"{label} must be a number")
    return number


def parse_date(value: Any, label: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = cell_text(value) or ""
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ImportRowError(f"{label} must be a date like 31/12/2026")


def row_to_product_data(values: List[Any], columns: Dict[str, int]) -> dict:
    """ProductCreate fields for one row; raises ImportRowError listing every problem found"""
    raw = {f: values[i] for f, i in columns.items()
           if values[i] is not None and str(values[i]).strip() != ""}
    data: dict = {}
    errors: List[str] = []

    def check(fn, *args):
        try:
            return fn(*args)
        except ImportRowError as exc:
            errors.append(str(exc))
            return None

    data["item_name"] = cell_text(raw.get("item_name"))
    if not data["item_name"]:
        errors.append("Item Name is required")

    type_text = re.sub(r"[^a-z]", "", str(raw.get("product_type") or "").lower())
    product_type = PRODUCT_TYPES.get(type_text)
    if not raw.get("product_type"):
        errors.append("Product Type is required (Manufacturer, Own packed or Loose)")
    elif not product_type:
        errors.append(f"Product Type '{raw['product_type']}' is not valid; use Manufacturer, Own packed or Loose")
    data["product_type"] = product_type

    for name in TEXT_FIELDS:
        if name in raw:
            data[name] = cell_text(raw[name])

    for name in DECIMAL_FIELDS:
        if name in raw:
            data[name] = check(parse_decimal, raw[name], LABELS[name])

    if "net_unit" in raw:
        data["net_unit"] = UNITS.get(str(raw["net_unit"]).strip().lower())
        if not data["net_unit"]:
            errors.append(f"Pack Unit '{raw['net_unit']}' is not valid; use g, kg, ml, L, pcs or pages")
    if product_type == Product.TYPE_LOOSE:
        data["unit_type"] = LOOSE_UNIT
    if "expiry_date" in raw:
        data["expiry_date"] = check(parse_date, raw["expiry_date"], LABELS["expiry_date"])
    data["reorder_level"] = IMPORT_REORDER_LEVEL

    if errors:
        raise ImportRowError("; ".join(errors))
    return {k: v for k, v in data.items() if v is not None}


def validation_message(exc: ValidationError) -> str:
    """Readable message from a ProductCreate validation error"""
    parts = []
    for err in exc.errors():
        msg = str(err.get("msg", "")).replace("Value error, ", "")
        label = LABELS.get(str(err["loc"][0])) if err.get("loc") else None
        parts.append(f"{label}: {msg}" if label else msg)
    return "; ".join(parts) or "Invalid row"


def build_import_template() -> bytes:
    """Excel template: a Products sheet with headers and dropdowns, and an Instructions sheet"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    bold, fill = Font(bold=True), PatternFill("solid", fgColor="DDEBF7")
    for col, (name, header, required, _desc, _aliases) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=col, value=f"{header}*" if required else header)
        cell.font, cell.fill = bold, fill
        letter = cell.column_letter
        sheet.column_dimensions[letter].width = max(12, len(header) + 4)
        if name == "hsn_code":
            for row in range(2, MAX_IMPORT_ROWS + 2):
                sheet.cell(row=row, column=col).number_format = "@"
        elif name == "expiry_date":
            for row in range(2, MAX_IMPORT_ROWS + 2):
                sheet.cell(row=row, column=col).number_format = "DD/MM/YYYY"
    sheet.freeze_panes = "A2"

    def add_list(field_name: str, options: str):
        letter = sheet.cell(row=1, column=[c[0] for c in COLUMNS].index(field_name) + 1).column_letter
        rule = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True)
        rule.add(f"{letter}2:{letter}{MAX_IMPORT_ROWS + 1}")
        sheet.add_data_validation(rule)

    add_list("product_type", "Manufacturer,Own packed,Loose")
    add_list("net_unit", "g,kg,ml,L,pcs,pages")

    guide = workbook.create_sheet("Instructions")
    guide.append(["Column", "Required", "Description"])
    for _name, header, required, desc, _aliases in COLUMNS:
        guide.append([header, "Yes" if required else "", desc])
    guide.append([])
    guide.append(["Fill the Products sheet, one product per row. Rows whose display name and type "
                  "already exist are skipped."])
    guide.append([])
    guide.append(["Examples"])
    guide.append([header for _n, header, *_ in COLUMNS])
    for example in EXAMPLES:
        guide.append(example)
    for cell in guide[1]:
        cell.font = bold
    guide.column_dimensions["A"].width = 18
    guide.column_dimensions["B"].width = 10
    guide.column_dimensions["C"].width = 90

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
